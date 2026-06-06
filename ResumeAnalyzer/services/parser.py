import re
import os
import spacy
import nltk
from nltk.tokenize import sent_tokenize, word_tokenize

try:
    nlp = spacy.load("en_core_web_sm")
except OSError:
    from spacy.cli import download
    download("en_core_web_sm")
    nlp = spacy.load("en_core_web_sm")

nltk_packages = {
    'punkt': 'tokenizers/punkt',
    'stopwords': 'corpora/stopwords',
    'averaged_perceptron_tagger': 'taggers/averaged_perceptron_tagger',
    'maxent_ne_chunker': 'chunkers/maxent_ne_chunker',
    'words': 'corpora/words'
}

for pkg, path in nltk_packages.items():
    try:
        nltk.data.find(path)
    except LookupError:
        try:
            nltk.download(pkg, quiet=True)
        except Exception:
            pass


class ResumeParser:
    def __init__(self):
        self.nlp = nlp

    def parse(self, file_path: str) -> dict:
        ext = os.path.splitext(file_path)[1].lower()
        if ext == '.pdf':
            text = self._extract_pdf(file_path)
        elif ext == '.docx':
            text = self._extract_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

        return self._extract_entities(text)

    def _extract_pdf(self, file_path: str) -> str:
        text = ""
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception:
            try:
                import PyPDF2
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    for page in reader.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
            except Exception as e:
                raise RuntimeError(f"PDF extraction failed: {str(e)}")
        return text.strip()

    def _extract_docx(self, file_path: str) -> str:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)

    def _extract_entities(self, text: str) -> dict:
        result = {
            'raw_text': text,
            'name': self._extract_name(text),
            'email': self._extract_email(text),
            'phone': self._extract_phone(text),
            'linkedin': self._extract_linkedin(text),
            'github': self._extract_github(text),
            'skills': self._extract_skills_section(text),
            'education': self._extract_education(text),
            'experience': self._extract_experience(text),
            'certifications': self._extract_certifications(text),
            'projects': self._extract_projects(text),
            'summary': self._extract_summary(text),
            'years_of_experience': self._calculate_years_experience(text),
        }
        return result

    def _extract_name(self, text: str) -> str:
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        doc = self.nlp(text[:500])
        for ent in doc.ents:
            if ent.label_ == "PERSON":
                return ent.text.strip()
        if lines:
            first_line = lines[0]
            if len(first_line.split()) <= 5 and not re.search(r'[@\d]', first_line):
                return first_line
        return "Not Found"

    def _extract_email(self, text: str) -> str:
        pattern = r'\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}\b'
        matches = re.findall(pattern, text)
        return matches[0] if matches else "Not Found"

    def _extract_phone(self, text: str) -> str:
        patterns = [
            r'\+?1?\s*[\-.]?\s*\(?\d{3}\)?[\s\-.]?\d{3}[\s\-.]?\d{4}',
            r'\+?\d{1,3}[\s\-.]?\d{3,5}[\s\-.]?\d{3,5}[\s\-.]?\d{3,5}',
        ]
        for pattern in patterns:
            matches = re.findall(pattern, text)
            if matches:
                phone = re.sub(r'[^\d+\-() ]', '', matches[0]).strip()
                if len(re.sub(r'\D', '', phone)) >= 10:
                    return phone
        return "Not Found"

    def _extract_linkedin(self, text: str) -> str:
        pattern = r'linkedin\.com/in/[\w\-]+'
        matches = re.findall(pattern, text, re.IGNORECASE)
        return matches[0] if matches else ""

    def _extract_github(self, text: str) -> str:
        pattern = r'github\.com/[\w\-]+'
        matches = re.findall(pattern, text, re.IGNORECASE)
        return matches[0] if matches else ""

    def _extract_skills_section(self, text: str) -> list:
        skills_section = self._get_section(text, [
            'skills', 'technical skills', 'core competencies',
            'technologies', 'tools', 'expertise', 'proficiencies'
        ])
        if not skills_section:
            return []

        skills = []
        lines = skills_section.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            parts = re.split(r'[,|•·;/]', line)
            for part in parts:
                skill = part.strip().strip('•-–').strip()
                if skill and 2 <= len(skill) <= 50 and not skill.isnumeric():
                    skills.append(skill)

        return list(dict.fromkeys([s for s in skills if s]))

    def _extract_education(self, text: str) -> list:
        section = self._get_section(text, ['education', 'academic background', 'qualifications', 'academic qualifications'])
        if not section:
            section = text

        education = []
        degree_patterns = [
            r'(Bachelor|Master|PhD|Ph\.D|B\.S|M\.S|B\.E|M\.E|B\.Tech|M\.Tech|MBA|BBA|Associate|Diploma|B\.Sc|M\.Sc|B\.A|M\.A)[\w\s,\.]*(?:in|of)?[\w\s,\.]{0,60}',
        ]
        for pattern in degree_patterns:
            matches = re.findall(pattern, section, re.IGNORECASE)
            for match in matches:
                degree_text = match.strip()
                if degree_text and degree_text not in education:
                    education.append(degree_text)

        university_patterns = [
            r'(?:University|College|Institute|School|Academy) of [\w\s,]+',
            r'[\w\s]+ (?:University|College|Institute|School)',
        ]
        doc = self.nlp(section[:2000])
        for ent in doc.ents:
            if ent.label_ == "ORG" and any(w in ent.text for w in ['University', 'College', 'Institute', 'School', 'Academy']):
                if ent.text not in education:
                    education.append(ent.text)

        return education[:8]

    def _extract_experience(self, text: str) -> list:
        section = self._get_section(text, [
            'experience', 'work experience', 'employment', 'professional experience',
            'work history', 'career history', 'professional background'
        ])
        if not section:
            return []

        experience = []
        lines = [l.strip() for l in section.split('\n') if l.strip()]
        date_pattern = r'\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|January|February|March|April|June|July|August|September|October|November|December)[\w\s,\-–]*\d{4}\b'
        year_pattern = r'\b(19|20)\d{2}\s*[\-–to]+\s*(19|20)\d{2}|present|current\b'

        i = 0
        current_entry = []
        while i < len(lines):
            line = lines[i]
            if re.search(date_pattern, line, re.IGNORECASE) or re.search(year_pattern, line, re.IGNORECASE):
                if current_entry:
                    experience.append(' | '.join(current_entry[:3]))
                current_entry = [line]
            elif current_entry and len(line) > 5:
                current_entry.append(line)
            i += 1

        if current_entry:
            experience.append(' | '.join(current_entry[:3]))

        if not experience:
            doc = self.nlp(section[:2000])
            for ent in doc.ents:
                if ent.label_ == "ORG":
                    experience.append(ent.text)

        return experience[:10]

    def _extract_certifications(self, text: str) -> list:
        section = self._get_section(text, [
            'certifications', 'certificates', 'credentials',
            'professional certifications', 'licenses', 'achievements'
        ])
        certs = []

        cert_keywords = [
            'AWS', 'Azure', 'GCP', 'Google', 'Microsoft', 'Cisco', 'CompTIA',
            'Certified', 'Certificate', 'Certification', 'CISSP', 'CEH', 'CISM',
            'PMP', 'Scrum', 'ITIL', 'Oracle', 'VMware', 'Red Hat', 'RHCE',
            'Six Sigma', 'ISO', 'CPA', 'CFA', 'CCNA', 'CCNP', 'CISA',
            'Kubernetes', 'Docker', 'Terraform', 'Security+', 'Network+'
        ]

        search_text = section if section else text
        lines = search_text.split('\n')
        for line in lines:
            line = line.strip()
            if any(kw.lower() in line.lower() for kw in cert_keywords):
                clean = line.strip('•-–*').strip()
                if clean and len(clean) > 5 and clean not in certs:
                    certs.append(clean)

        return certs[:10]

    def _extract_projects(self, text: str) -> list:
        section = self._get_section(text, [
            'projects', 'personal projects', 'academic projects',
            'side projects', 'portfolio', 'key projects'
        ])
        if not section:
            return []

        projects = []
        lines = [l.strip() for l in section.split('\n') if l.strip()]
        for line in lines:
            if len(line) > 10 and not line.startswith(('•', '-', '*', '–')):
                if not re.match(r'^[\d\.\s]+$', line):
                    projects.append(line)
            elif line.startswith(('•', '-', '*', '–')):
                clean = line.strip('•-*–').strip()
                if len(clean) > 10:
                    projects.append(clean)

        return projects[:8]

    def _extract_summary(self, text: str) -> str:
        section = self._get_section(text, [
            'summary', 'objective', 'profile', 'about me',
            'professional summary', 'career objective', 'overview'
        ])
        if section:
            lines = [l.strip() for l in section.split('\n') if l.strip()]
            return ' '.join(lines[:3])[:500]
        return ""

    def _calculate_years_experience(self, text: str) -> float:
        year_ranges = re.findall(
            r'\b(20\d{2}|19\d{2})\s*[\-–to]+\s*(20\d{2}|19\d{2}|present|current|now)\b',
            text, re.IGNORECASE
        )
        import datetime
        current_year = datetime.datetime.now().year
        total_years = 0.0

        for start, end in year_ranges:
            try:
                start_year = int(start)
                end_year = current_year if end.lower() in ['present', 'current', 'now'] else int(end)
                diff = end_year - start_year
                if 0 < diff <= 50:
                    total_years += diff
            except ValueError:
                continue

        if total_years == 0:
            exp_match = re.search(r'(\d+)\+?\s*years?\s*(?:of\s+)?(?:experience|exp)', text, re.IGNORECASE)
            if exp_match:
                total_years = float(exp_match.group(1))

        return round(min(total_years, 40), 1)

    def _get_section(self, text: str, section_names: list) -> str:
        lines = text.split('\n')
        section_headers = [
            'experience', 'education', 'skills', 'certifications', 'projects',
            'summary', 'objective', 'achievements', 'awards', 'references',
            'publications', 'languages', 'interests', 'hobbies', 'contact'
        ]

        start_idx = -1
        end_idx = len(lines)

        for i, line in enumerate(lines):
            line_lower = line.strip().lower()
            for name in section_names:
                if line_lower == name or line_lower.startswith(name + ':') or line_lower.startswith(name + ' '):
                    if len(line_lower) <= len(name) + 10:
                        start_idx = i + 1
                        break
            if start_idx != -1:
                break

        if start_idx == -1:
            return ""

        for j in range(start_idx, len(lines)):
            line_lower = lines[j].strip().lower()
            if j == start_idx:
                continue
            for header in section_headers:
                is_current_section = any(name in line_lower for name in section_names)
                if not is_current_section and line_lower == header or line_lower.startswith(header + ':'):
                    if len(line_lower) <= len(header) + 10:
                        end_idx = j
                        break
            if end_idx != len(lines):
                break

        return '\n'.join(lines[start_idx:end_idx]).strip()
